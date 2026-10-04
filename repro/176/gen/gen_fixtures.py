#!/usr/bin/env python3
# Generates all EP02 quick-test fixtures (fictional, self-made) + truth files. Deterministic.
# usage: .venv-gen/bin/python gen/gen_fixtures.py   (from the episode dir)
import csv, json, os, random, shutil, subprocess, tempfile
from pathlib import Path
H = Path(__file__).resolve().parent.parent
FX, TR = H / "fixtures", H / "truth"
for d in (FX, TR): d.mkdir(exist_ok=True)
for lv in ("L1", "L2", "L3"): (FX / lv).mkdir(exist_ok=True)
rnd = random.Random(20261004)
CSS = "<style>body{font-family:'DejaVu Sans';font-size:10pt} table{border-collapse:collapse} td,th{border:1px solid #000;padding:3px}</style>"

def html2(fmt, html, out):
    tmp = Path(tempfile.mkdtemp(prefix="ep02gen-"))
    src = tmp / (out.stem + ".html"); src.write_text("<html><head><meta charset='utf-8'>" + CSS + "</head><body>" + html + "</body></html>", encoding="utf-8")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", fmt, "--outdir", str(tmp), str(src)], check=True, capture_output=True, timeout=180)
    shutil.move(str(tmp / (out.stem + "." + fmt.split(":")[0])), out); shutil.rmtree(tmp)

def table(head, rows):
    h = "<tr>" + "".join(f"<th>{c}</th>" for c in head) + "</tr>"
    return "<table>" + h + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows) + "</table>"
PB = "<p style='page-break-before: always'></p>"

# ---------------- L1: one PDF, extract Table 2 (12 rows), ignore Table 1
desc = ["Servo bracket", "Gear shaft 8 mm", "Bearing cap", "Cable harness", "Sensor mount", "Drive belt", "Motor plate", "Spacer kit", "Hinge pin", "Wheel hub", "Cover panel", "Fuse holder"]
units = ["pcs", "pcs", "pcs", "set", "pcs", "m", "pcs", "set", "pcs", "pcs", "pcs", "pcs"]
L1 = []
for i, (d, u) in enumerate(zip(desc, units)):
    L1.append({"part_no": f"LF-P-{101 + i * 7:04d}", "description": d, "unit": u,
               "unit_price_cny": round(rnd.randint(350, 98000) / 100, 2), "stock": rnd.randint(0, 900)})
contacts = [["Mira Okafor", "Purchasing lead", "4102"], ["Tomas Vey", "Quality", "4117"], ["Lin Arden", "Logistics", "4120"]]
html = ("<h2>Lanternfish Robotics (fictional company) - Supplier Report Q3 2026</h2>"
        "<p>This report is fictional test data made for an automated evaluation. Section 1 lists contacts; Section 2 lists the current parts price list.</p>"
        "<h3>Table 1: Contacts</h3>" + table(["Name", "Role", "Ext."], contacts) +
        "<p>Prices exclude VAT. Stock counted on 2026-09-30.</p>" + PB +
        "<h3>Table 2: Parts price list</h3>" +
        table(["Part No.", "Description", "Unit", "Unit Price (CNY)", "Stock"],
              [[r["part_no"], r["description"], r["unit"], f"{r['unit_price_cny']:,.2f}", r["stock"]] for r in L1]) +
        "<p>End of report.</p>")
html2("pdf", html, FX / "L1" / "supplier_report.pdf")

# ---------------- L2: three formats -> one CSV
MON = ["2026-07", "2026-08", "2026-09"]; SKU = ["AQ-100", "AQ-200", "AQ-300"]
L2 = []
for reg in ("north", "south", "east"):
    for m in MON:
        for s in SKU:
            u = rnd.randint(40, 600); L2.append({"region": reg, "month": m, "sku": s, "units": u, "revenue_cny": rnd.randint(30, 2400) * 100})
def rows(reg): return [r for r in L2 if r["region"] == reg]
EN = {"2026-07": "Jul 2026", "2026-08": "Aug 2026", "2026-09": "Sep 2026"}
html = ("<h2>North Region - Q3 2026 Sales (fictional)</h2><p>Revenue is stated in units of CNY 10,000 (wan yuan).</p>" +
        table(["Month", "SKU", "Units", "Revenue (CNY 10k)"], [[EN[r["month"]], r["sku"], r["units"], f"{r['revenue_cny'] / 10000:.2f}"] for r in rows("north")]))
html2("pdf", html, FX / "L2" / "north_q3.pdf")
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Notes"
ws.append(["说明 / Notes"]); ws.append(["虚构测试数据。销售数据在 Sales 表。Fictional test data; see sheet Sales."])
ws.append(["Old draft figures (do not use):"]); ws.append(["AQ-100", 999, 999999])
ws2 = wb.create_sheet("Sales"); ws2.append(["月份", "产品编号", "数量", "收入（元）"])
for r in rows("south"): ws2.append([f"2026年{int(r['month'][5:])}月", r["sku"], r["units"], r["revenue_cny"]])
ws2.append(["合计", "", sum(r["units"] for r in rows("south")), sum(r["revenue_cny"] for r in rows("south"))])
wb.save(FX / "L2" / "south_q3.xlsx")
from pptx import Presentation
from pptx.util import Inches
p = Presentation(); s1 = p.slides.add_slide(p.slide_layouts[0]); s1.shapes.title.text = "East Region Q3 2026 (fictional)"; s1.placeholders[1].text = "Revenue in kCNY (thousand yuan)"
s2 = p.slides.add_slide(p.slide_layouts[5]); s2.shapes.title.text = "Q3 sales by month"
er = rows("east"); tb = s2.shapes.add_table(len(er) + 1, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(5)).table
for j, c in enumerate(["Month", "Item", "Qty", "Revenue (kCNY)"]): tb.cell(0, j).text = c
for i, r in enumerate(er, 1):
    for j, v in enumerate([r["month"].replace("2026-", "") + "/2026", r["sku"], str(r["units"]), f"{r['revenue_cny'] / 1000:.1f}"]): tb.cell(i, j).text = v
p.save(FX / "L2" / "east_q3.pptx")

# ---------------- L3: cross-page register with merged vendor cells, running headers/footers, subtotals, footnote, amendments docx
V = ["Harborlight Logistics", "Quillstone Media", "Ferncliff Systems", "Tidewater Analytics", "Brightmoor Supplies"]
C = []
for i in range(24):
    mth, day = rnd.randint(1, 9), rnd.randint(13, 28)
    C.append({"contract_id": f"C-{1001 + i}", "vendor": V[min(i // 5, 4)], "start_date": f"2026-{mth:02d}-{day:02d}", "amount_cny": round(rnd.randint(500000, 9000000) / 100, 2)})
pages = [C[0:8], C[8:16], C[16:24]]
html = ""
for pi, pg in enumerate(pages):
    if pi: html += PB
    html += f"<p>Fictional Co. Internal - Contract Register - Page {pi + 1} of 3</p>"
    date_hdr = "Start (DD.MM.YYYY)" if pi == 2 else "Start (YYYY/MM/DD)"
    h = "<tr><th>Contract</th><th>Vendor</th><th>" + date_hdr + "</th><th>Amount</th></tr>"
    body = ""; k = 0
    while k < len(pg):
        n = 1
        while k + n < len(pg) and pg[k + n]["vendor"] == pg[k]["vendor"]: n += 1
        for j in range(n):
            r = pg[k + j]; y, m, d = r["start_date"].split("-")
            ds = f"{d}.{m}.{y}" if pi == 2 else f"{y}/{m}/{d}"
            vc = f"<td rowspan='{n}'>{r['vendor']}</td>" if j == 0 else ""
            body += f"<tr><td>{r['contract_id']}</td>{vc}<td>{ds}</td><td>CNY {r['amount_cny']:,.2f}</td></tr>"
        k += n
    sub = sum(r["amount_cny"] for r in pg)
    body += f"<tr><td>Page subtotal</td><td></td><td></td><td>CNY {sub:,.2f}</td></tr>"
    html += "<table>" + h + body + "</table>"
    if pi == 2: html += "<p>* Contract C-1017 was cancelled on 2026-09-12. Exclude it from the register.</p>"
    html += "<p>Draft v2 - not for distribution</p>"
html2("pdf", html, FX / "L3" / "contract_register.pdf")
from docx import Document
doc = Document(); doc.add_heading("Amendments to the contract register (fictional)", 1)
doc.add_paragraph("Apply these amendments on top of the register. Values in the New column replace the register values.")
AM = [("C-1004", "Amount", "CNY 61,250.00"), ("C-1011", "Amount", "CNY 18,900.50"), ("C-1020", "Start date", "2026-10-15"), ("C-1017", "Amount", "CNY 5,000.00")]
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
for j, c in enumerate(["Contract", "Field", "Old", "New"]): t.rows[0].cells[j].text = c
byid = {r["contract_id"]: r for r in C}
for cid, f, new in AM:
    r = byid[cid]; old = f"CNY {r['amount_cny']:,.2f}" if f == "Amount" else r["start_date"]
    cells = t.add_row().cells
    for j, v in enumerate([cid, f, old, new]): cells[j].text = v
doc.save(FX / "L3" / "amendments.docx")
final = []
for r in C:
    r = dict(r)
    if r["contract_id"] == "C-1017": continue
    for cid, f, new in AM:
        if cid == r["contract_id"]:
            if f == "Amount": r["amount_cny"] = float(new.replace("CNY ", "").replace(",", ""))
            else: r["start_date"] = new
    final.append(r)

def wcsv(path, head, data):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=head); w.writeheader(); [w.writerow({k: r[k] for k in head}) for r in data]
wcsv(TR / "L1_parts.csv", ["part_no", "description", "unit", "unit_price_cny", "stock"], L1)
wcsv(TR / "L2_q3_sales.csv", ["region", "month", "sku", "units", "revenue_cny"], L2)
json.dump({"rows": len(L2), "total_revenue_cny": sum(r["revenue_cny"] for r in L2)}, open(TR / "L2_summary.json", "w"), indent=1)
wcsv(TR / "L3_contracts.csv", ["contract_id", "vendor", "start_date", "amount_cny"], final)
print("L1", len(L1), "L2", len(L2), "L3", len(final))
